# Usage: run `python generate_doc.py` to create the word document (requires pip install python-docx)

import os
try:
    from docx import Document
except ImportError:
    print("Please install python-docx: pip install python-docx")
    exit(1)

def convert_md_to_docx(input_file, output_file):
    doc = Document()

    with open(input_file, 'r', encoding='utf-8') as f:
        for line in f:
            line = line.strip()
            if not line:
                continue

            if line.startswith('# '):
                doc.add_heading(line[2:], level=1)
            elif line.startswith('## '):
                doc.add_heading(line[3:], level=2)
            elif line.startswith('### '):
                doc.add_heading(line[4:], level=3)
            elif line.startswith('- '):
                doc.add_paragraph(line[2:], style='List Bullet')
            elif line.startswith('```'):
                continue
            else:
                doc.add_paragraph(line)

    doc.save(output_file)
    print(f"Success! {output_file} created.")

if __name__ == "__main__":
    if os.path.exists("Report_Project2.md"):
        convert_md_to_docx("Report_Project2.md", "Report_Project2.docx")
