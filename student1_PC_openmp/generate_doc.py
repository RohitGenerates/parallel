# Note to user: If you want to convert the Markdown to docx automatically, you can run this script.
# Make sure to install python-docx first: pip install python-docx

import os
try:
    from docx import Document
except ImportError:
    print("Please run `pip install python-docx` first.")
    exit(1)

def build_docx(md_filename, docx_filename):
    doc = Document()
    doc.add_heading(md_filename.replace('.md', ''), 0)

    with open(md_filename, 'r', encoding='utf-8') as f:
        for line in f:
            line = line.strip()
            if not line:
                continue

            if line.startswith('# '):
                doc.add_heading(line[2:], level=1)
            elif line.startswith('## '):
                doc.add_heading(line[3:], level=2)
            elif line.startswith('### '):
                doc.add_heading(line[4:], level=3)
            elif line.startswith('- '):
                doc.add_paragraph(line[2:], style='List Bullet')
            else:
                doc.add_paragraph(line)

    doc.save(docx_filename)
    print(f"Generated {docx_filename}")

if __name__ == "__main__":
    if os.path.exists("Report_Project1.md"):
        build_docx("Report_Project1.md", "Report_Project1.docx")
